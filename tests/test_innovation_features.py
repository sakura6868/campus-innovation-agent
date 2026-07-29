from __future__ import annotations

import sys
import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4


PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

import db  # noqa: E402
import api  # noqa: E402
from parsing.pdf_extractor import parse_pdf  # noqa: E402
from portfolio.optimizer import optimize_portfolios  # noqa: E402
from radar.service import _normalise, analyse_change, run_watch, validate_public_url  # noqa: E402
from schemas import (  # noqa: E402
    Citation,
    Competition,
    CompetitionCategory,
    DataStatus,
    EducationLevel,
    Grade,
    PortfolioPreferences,
    TrustedLevel,
    UserProfile,
)
from trust import REQUIRED_EVIDENCE_FIELDS  # noqa: E402


def _minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 14 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


def _ready_comp(index: int, deadline_days: int, skills: list[str]) -> Competition:
    checked = date.today().isoformat()
    url = f"https://contest.example.edu/official-{index}.pdf"
    return Competition(
        competition_id=f"portfolio_{index}",
        competition_name=f"组合测试赛事{index}",
        document_year=date.today().year,
        category=(CompetitionCategory.SOFTWARE if index % 2 else CompetitionCategory.MODELING),
        eligible_students=[EducationLevel.UNDERGRADUATE],
        registration_deadline=date.today() + timedelta(days=deadline_days),
        required_skills=skills,
        required_materials=["报名表", "作品"],
        official_source_url=url,
        source_acquired_date=checked,
        official_source_status="found",
        last_verified_at=checked,
        data_status=DataStatus.VERIFIED,
        trusted_level=TrustedLevel.A,
        evidence=[
            Citation(
                field=field,
                page=1,
                source_text=f"官方原文 {field}",
                document_name=f"official-{index}.pdf",
                source_url=url,
                acquired_date=checked,
                last_verified_at=checked,
                trusted_level=TrustedLevel.A,
            )
            for field in REQUIRED_EVIDENCE_FIELDS
        ],
    )


class PortfolioOptimizerTests(unittest.TestCase):
    def test_three_modes_respect_capacity_and_trust_gate(self) -> None:
        user = UserProfile(
            user_id="portfolio-user",
            education_level=EducationLevel.UNDERGRADUATE,
            grade=Grade.SOPHOMORE,
            major="软件工程",
            skills=["Python", "前端"],
            weekly_available_hours=18,
            expected_team_size=3,
            privacy_consent=True,
        )
        competitions = [
            _ready_comp(1, 60, ["Python", "前端"]),
            _ready_comp(2, 75, ["数学建模", "Python"]),
            _ready_comp(3, 90, ["产品设计"]),
            _ready_comp(4, 105, ["数据分析"]),
        ]
        unverified = _ready_comp(9, 80, ["Python"]).model_copy(
            update={"data_status": DataStatus.UNVERIFIED, "trusted_level": TrustedLevel.B}
        )
        result = optimize_portfolios(
            user,
            competitions + [unverified],
            [],
            PortfolioPreferences(max_competitions=3, horizon_weeks=16),
        )
        self.assertEqual([plan.mode for plan in result.plans], ["steady", "balanced", "sprint"])
        for plan in result.plans:
            self.assertLessEqual(plan.peak_weekly_load, plan.capacity_hours + 0.01)
            self.assertNotIn("portfolio_9", {item.competition_id for item in plan.items})


class PortfolioApplyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        db.init_db()

    def test_apply_is_atomic_and_idempotent(self) -> None:
        suffix = uuid4().hex
        comps = [
            _ready_comp(201, 90, ["Python"]).model_copy(
                update={"competition_id": f"atomic_a_{suffix}"}
            ),
            _ready_comp(202, 100, ["产品设计"]).model_copy(
                update={"competition_id": f"atomic_b_{suffix}"}
            ),
        ]
        for comp in comps:
            db.upsert_competition(comp)
        self.addCleanup(self._cleanup_competitions, [comp.competition_id for comp in comps])

        projects, failures = db.create_user_projects_atomic(
            "test", [comps[0].competition_id, f"missing_{suffix}"]
        )
        self.assertEqual(projects, [])
        self.assertEqual(len(failures), 1)
        self.assertNotIn(
            comps[0].competition_id,
            {item.competition_id for item in db.list_user_projects("test")},
        )

        first, failures = db.create_user_projects_atomic(
            "test", [comp.competition_id for comp in comps]
        )
        second, failures_again = db.create_user_projects_atomic(
            "test", [comp.competition_id for comp in comps]
        )
        self.assertFalse(failures)
        self.assertFalse(failures_again)
        self.assertEqual([item.project_id for item in first], [item.project_id for item in second])

    @staticmethod
    def _cleanup_competitions(competition_ids: list[str]) -> None:
        with db.session_scope() as session:
            project_ids = [
                row[0]
                for row in session.query(db.UserProjectModel.project_id)
                .filter(db.UserProjectModel.competition_id.in_(competition_ids))
                .all()
            ]
            if project_ids:
                session.query(db.ProjectItemModel).filter(
                    db.ProjectItemModel.project_id.in_(project_ids)
                ).delete(synchronize_session=False)
            session.query(db.UserProjectModel).filter(
                db.UserProjectModel.competition_id.in_(competition_ids)
            ).delete(synchronize_session=False)
            session.query(db.CitationModel).filter(
                db.CitationModel.competition_id.in_(competition_ids)
            ).delete(synchronize_session=False)
            session.query(db.CompetitionModel).filter(
                db.CompetitionModel.competition_id.in_(competition_ids)
            ).delete(synchronize_session=False)


class RadarTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        db.init_db()

    def test_critical_deadline_change_is_proposed_not_auto_applied(self) -> None:
        old = "赛事通知\n报名截止时间：2026年09月30日\n参赛对象：本科生"
        new = "赛事通知（更新）\n报名截止时间调整为：2026年09月20日\n参赛对象：本科生"
        analysis = analyse_change(old, new)
        self.assertEqual(analysis["severity"], "high")
        self.assertEqual(
            analysis["proposed_changes"]["registration_deadline"]["value"], "2026-09-20"
        )

    def test_demo_scan_creates_baseline_then_pending_change(self) -> None:
        watch = db.create_source_watch(
            "china_softcup_2026",
            f"https://example.edu/notice-{uuid4().hex}",
        )
        self.addCleanup(self._cleanup_watch, watch["watch_id"])
        first = run_watch(watch["watch_id"], "报名截止时间：2026年09月30日")
        second = run_watch(watch["watch_id"], "报名截止时间：2026年09月20日")
        self.assertEqual(first["status"], "baseline_created")
        self.assertEqual(second["status"], "change_detected")
        self.assertEqual(second["severity"], "high")
        event = next(item for item in db.list_change_events() if item["event_id"] == second["event_id"])
        self.assertEqual(event["status"], "pending")

    @staticmethod
    def _cleanup_watch(watch_id: int) -> None:
        with db.session_scope() as session:
            session.query(db.SourceChangeEventModel).filter_by(watch_id=watch_id).delete()
            session.query(db.SourceSnapshotModel).filter_by(watch_id=watch_id).delete()
            session.query(db.SourceWatchModel).filter_by(watch_id=watch_id).delete()

    def test_private_source_urls_are_rejected(self) -> None:
        for url in ("http://127.0.0.1/admin", "http://localhost/internal", "file:///tmp/a"):
            with self.subTest(url=url), self.assertRaises(ValueError):
                validate_public_url(url)

    def test_docx_source_is_detected_and_normalized(self) -> None:
        from docx import Document

        document = Document()
        document.add_heading("赛事官方通知", level=1)
        document.add_paragraph("报名截止时间：2026年09月30日")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "参赛对象"
        table.cell(0, 1).text = "在校本科生"
        buffer = BytesIO()
        document.save(buffer)

        normalized, kind = _normalise(
            buffer.getvalue(),
            "application/octet-stream",
            "auto",
            None,
        )
        self.assertEqual(kind, "docx")
        self.assertIn("报名截止时间：2026年09月30日", normalized)
        self.assertIn("参赛对象 | 在校本科生", normalized)


if __name__ == "__main__":
    unittest.main()
